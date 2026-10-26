"""Create the private data/ folder: tracker, JD archive, progress log, notes.

  python scripts/setup_workspace.py           # keeps existing files
  python scripts/setup_workspace.py --force   # recreate the tracker (asks no questions!)
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from _common import (COLUMNS, DATA, JD_ARCHIVE, NOTES, PROGRESS, ROOT, STATUSES, TMP_JD, TRACKER,  # noqa: E402
                     need, today)


def make_tracker():
    openpyxl = need("openpyxl")
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws["A1"] = "Job application tracker · 投递记录"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(COLUMNS)
    fill = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    widths = [5, 34, 22, 14, 12, 12, 8, 8, 24, 20, 30, 34, 60, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("N3:N5000")

    sm = wb.create_sheet("Summary")
    sm.append(["Status", "Count"])
    for i, s in enumerate(STATUSES, start=2):
        sm.append([s, f"=COUNTIF(Applications!N3:N5000,A{i})"])
    sm.append(["Total 总数", "=COUNTA(Applications!A3:A5000)"])
    for c in sm[1]:
        c.font = Font(bold=True)
    sm.column_dimensions["A"].width = 24
    wb.save(TRACKER)


def write_if_missing(path: Path, text: str):
    if path.exists():
        print(f"已存在，跳过 / exists, kept: {path.relative_to(ROOT) if ROOT in path.parents else path}")
        return
    path.write_text(text, encoding="utf-8")
    print(f"已创建 / created: {path.relative_to(ROOT) if ROOT in path.parents else path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="recreate tracker even if it exists")
    a = ap.parse_args()
    DATA.mkdir(parents=True, exist_ok=True)
    (DATA / "cover_letters").mkdir(exist_ok=True)
    (DATA / "interview_prep").mkdir(exist_ok=True)
    TMP_JD.mkdir(exist_ok=True)

    if TRACKER.exists() and not a.force:
        print(f"已存在，跳过 / exists, kept: {TRACKER}")
    else:
        make_tracker()
        print(f"已创建 / created: {TRACKER}")

    tpl = ROOT / "templates"
    write_if_missing(JD_ARCHIVE, (tpl / "JD_archive.md").read_text(encoding="utf-8"))
    write_if_missing(PROGRESS, (tpl / "progress.md").read_text(encoding="utf-8").replace("{{DATE}}", today()))
    write_if_missing(NOTES, (tpl / "notes.md").read_text(encoding="utf-8"))
    print("\n✓ 工作区就绪 / workspace ready. 下一步 / next: python scripts/make_cheatsheet.py")


if __name__ == "__main__":
    main()
