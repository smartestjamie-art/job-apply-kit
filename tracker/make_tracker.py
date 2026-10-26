"""Create a blank tracker: python tracker/make_tracker.py [out.xlsx]"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNS = ["#", "Job Title", "Company", "Location", "Date Applied", "Salary", "Headcount",
           "Fit Score", "Cover Letter", "CV Version", "Job Source", "Flow", "Notes", "Status"]
STATUSES = ["Applied 已投/等待", "Assessment 测评中", "Interview 面试中", "Offer", "Rejected 被拒", "Withdrawn 撤回"]

out = sys.argv[1] if len(sys.argv) > 1 else "tracker/applications.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Applications"
ws.append(["Job application tracker"])
ws.append(COLUMNS)
head = PatternFill("solid", start_color="1F3864", end_color="1F3864")
for c in ws[2]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head
    c.alignment = Alignment(horizontal="center", vertical="center")
for col, w in zip("ABCDEFGHIJKLMN", [5, 34, 22, 16, 13, 12, 10, 9, 10, 24, 24, 36, 60, 20]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A3"

sm = wb.create_sheet("Summary")
sm.append(["Status", "Count"])
for i, s in enumerate(STATUSES, start=2):
    sm.append([s, f"=COUNTIF(Applications!N3:N1000,A{i})"])
sm.append(["Total", "=COUNTA(Applications!A3:A1000)"])
wb.save(out)
print("created", out)
