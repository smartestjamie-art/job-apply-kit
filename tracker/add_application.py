"""Append one application row, then re-open the file to verify it was saved.

Example:
  python tracker/add_application.py tracker/applications.xlsx \
    --title "Graduate Engineer" --company "Acme" --location "London" --date 2026-01-31 \
    --fit 4 --cover Y --cv "primary/Your Name.pdf" --source "Company site" \
    --flow "Applied" --notes "..." --status "Applied 已投/等待"
"""
import argparse
import copy
import openpyxl

p = argparse.ArgumentParser()
p.add_argument("path")
for k in ["title", "company", "location", "date", "salary", "headcount", "fit", "cover", "cv", "source", "flow", "notes", "status"]:
    p.add_argument("--" + k, default="")
a = p.parse_args()

wb = openpyxl.load_workbook(a.path)
ws = wb["Applications"]
last = ws.max_row
while last > 2 and ws.cell(last, 1).value in (None, ""):
    last -= 1
num = (ws.cell(last, 1).value or 0) + 1 if last > 2 else 1
row = last + 1
for c in range(1, 15):                       # copy the style of the previous row
    ws.cell(row, c)._style = copy.copy(ws.cell(last, c)._style)
vals = [num, a.title, a.company, a.location, a.date, a.salary or "Not listed", a.headcount or None,
        a.fit, a.cover, a.cv, a.source, a.flow, a.notes, a.status or "Applied 已投/等待"]
for i, v in enumerate(vals, 1):
    ws.cell(row, i).value = v

# widen Summary COUNTIF ranges if you outgrow N1000 (default is already large)
wb.save(a.path)

chk = openpyxl.load_workbook(a.path)["Applications"]
assert chk.cell(row, 2).value == a.title, "verification failed: row not saved"
print(f"added #{num} at row {row}")
